"""Generate the Sigurado project Gantt chart as a landscape short-bond XLSX.

Schedule: Mon 7:30-10:30 AM, Wed 7:30-10:30 AM, Sun 8:00 AM-5:00 PM.
Coverage: 20 July 2026 through 14 October 2026 (final defense).

Run:  python generate_gantt.py
"""

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

START = date(2026, 7, 20)   # Monday
END = date(2026, 10, 14)    # Wednesday, final defense
OUTFILE = Path(__file__).with_name("Sigurado_Gantt_Chart.xlsx")

FONT = "Arial"

# weekday -> (label, hours, time range)
SESSION_RULES = {
    0: ("M", 3.0, "7:30-10:30 AM"),
    2: ("W", 3.0, "7:30-10:30 AM"),
    6: ("Su", 9.0, "8:00 AM-5:00 PM"),
}

PHASES = {
    "P1": ("Planning and Feasibility", "1F4E79"),
    "P2": ("Requirements Analysis", "2E75B6"),
    "P3": ("System Design", "5B9BD5"),
    "P4": ("Implementation", "70AD47"),
    "P5": ("Testing (V and V)", "ED7D31"),
    "P6": ("Documentation and Defense", "7030A0"),
}

# (id, phase, task name, start, end)
TASKS = [
    (1, "P1", "Project kickoff, team roles and RACI", date(2026, 7, 20), date(2026, 7, 22)),
    (2, "P1", "Problem definition and feasibility study", date(2026, 7, 20), date(2026, 7, 26)),
    (3, "P1", "Project charter and scope sign-off", date(2026, 7, 22), date(2026, 7, 26)),
    (4, "P2", "Faculty interview", date(2026, 7, 26), date(2026, 8, 2)),
    (5, "P2", "Functional and non-functional requirements", date(2026, 7, 27), date(2026, 8, 2)),
    (6, "P2", "SRS document and approval", date(2026, 8, 2), date(2026, 8, 9)),
    (7, "P3", "System architecture and data flow diagrams", date(2026, 8, 3), date(2026, 8, 9)),
    (8, "P3", "Database design (ERD) and schema review", date(2026, 8, 3), date(2026, 8, 9)),
    (9, "P3", "Hardware schematic and wiring diagram", date(2026, 8, 5), date(2026, 8, 16)),
    (10, "P3", "UI and UX wireframes, design review", date(2026, 8, 9), date(2026, 8, 16)),
    (11, "P3", "SDD document and sign-off", date(2026, 8, 10), date(2026, 8, 16)),
    (12, "P4", "Source suppliers and finalize BOM", date(2026, 7, 27), date(2026, 8, 2)),
    (13, "P4", "Order hardware components", date(2026, 8, 3), date(2026, 8, 5)),
    (14, "P4", "Receive and inspect components (lead time)", date(2026, 8, 10), date(2026, 8, 23)),
    (15, "P4", "Backend refinement and security hardening", date(2026, 8, 17), date(2026, 9, 6)),
    (16, "P4", "Frontend refinement and dashboard polish", date(2026, 8, 17), date(2026, 9, 6)),
    (17, "P4", "Firmware: fingerprint sensor driver", date(2026, 8, 24), date(2026, 9, 6)),
    (18, "P4", "Firmware: I2C LCD, buttons, Wi-Fi, NTP", date(2026, 8, 24), date(2026, 9, 6)),
    (19, "P4", "Firmware: device API client and enrollment", date(2026, 8, 31), date(2026, 9, 13)),
    (20, "P4", "Circuit assembly: door node and relay lock", date(2026, 9, 7), date(2026, 9, 20)),
    (21, "P4", "Circuit assembly: box node, solenoid driver", date(2026, 9, 7), date(2026, 9, 20)),
    (22, "P4", "Enclosure fabrication and mounting", date(2026, 9, 14), date(2026, 9, 27)),
    (23, "P4", "Hardware and software integration", date(2026, 9, 14), date(2026, 9, 27)),
    (24, "P5", "Unit and integration testing", date(2026, 9, 7), date(2026, 9, 20)),
    (25, "P5", "System testing: door to box sequential gate", date(2026, 9, 21), date(2026, 10, 4)),
    (26, "P5", "Security testing: tailgating, expired session", date(2026, 9, 21), date(2026, 10, 4)),
    (27, "P5", "Defect fixing and regression retest", date(2026, 9, 28), date(2026, 10, 4)),
    (28, "P6", "User manual and technical documentation", date(2026, 9, 28), date(2026, 10, 7)),
    (29, "P6", "Defense preparation, dry run and rehearsal", date(2026, 10, 8), date(2026, 10, 12)),
    (30, "P6", "FINAL DEFENSE and system demonstration", date(2026, 10, 14), date(2026, 10, 14)),
]

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_sessions():
    """Every working session in range, tagged with its week number."""
    sessions = []
    day = START
    while day <= END:
        rule = SESSION_RULES.get(day.weekday())
        if rule:
            label, hours, _ = rule
            week = (day - START).days // 7 + 1
            sessions.append({"date": day, "label": label, "hours": hours, "week": week})
        day += timedelta(days=1)
    return sessions


def styled(cell, *, bold=False, size=8, color="000000", fill=None,
           align="center", wrap=False, border=True):
    cell.font = Font(name=FONT, bold=bold, size=size, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if border:
        cell.border = BORDER
    return cell


def main():
    sessions = build_sessions()
    weeks = sorted({s["week"] for s in sessions})
    total_hours = sum(s["hours"] for s in sessions)

    wb = Workbook()
    ws = wb.active
    ws.title = "Gantt Chart"

    first_session_col = 7           # columns A-F hold task metadata
    last_col = first_session_col + len(sessions) - 1
    last_letter = get_column_letter(last_col)

    # ---------- title block ----------
    ws.merge_cells(f"A1:{last_letter}1")
    styled(ws["A1"], bold=True, size=14, align="left", border=False)
    ws["A1"] = "Sigurado: Dual-Biometric Sequential Access and Materials Accountability System"

    ws.merge_cells(f"A2:{last_letter}2")
    styled(ws["A2"], size=8, align="left", border=False, color="404040")
    ws["A2"] = (
        f"Project Gantt Chart for {START.strftime('%B %d, %Y')} to {END.strftime('%B %d, %Y')}.   "
        f"{len(weeks)} weeks, {len(sessions)} work sessions, {total_hours:.0f} total hours.   "
        "Sessions: Monday 7:30-10:30 AM (3 hours), Wednesday 7:30-10:30 AM (3 hours), "
        "Sunday 8:00 AM-5:00 PM (9 hours)."
    )
    ws.row_dimensions[1].height = 19
    ws.row_dimensions[2].height = 13
    ws.row_dimensions[3].height = 4

    # ---------- header rows ----------
    head_row, sub_row = 4, 5
    headers = [("ID", 5), ("Task", 41), ("Phase", 7), ("Start", 10), ("End", 10), ("Hours", 7)]
    for idx, (label, width) in enumerate(headers, start=1):
        col = get_column_letter(idx)
        ws.column_dimensions[col].width = width
        ws.merge_cells(f"{col}{head_row}:{col}{sub_row}")
        styled(ws[f"{col}{head_row}"], bold=True, size=8, color="FFFFFF", fill="1F4E79")

    for idx, (label, _) in enumerate(headers, start=1):
        ws.cell(row=head_row, column=idx).value = label

    # week group headers spanning their sessions
    col = first_session_col
    for week in weeks:
        members = [s for s in sessions if s["week"] == week]
        span = len(members)
        start_letter = get_column_letter(col)
        end_letter = get_column_letter(col + span - 1)
        ws.merge_cells(f"{start_letter}{head_row}:{end_letter}{head_row}")
        cell = ws[f"{start_letter}{head_row}"]
        styled(cell, bold=True, size=7, color="FFFFFF", fill="2E75B6")
        cell.value = f"W{week}"

        for offset, session in enumerate(members):
            c = ws.cell(row=sub_row, column=col + offset)
            is_sun = session["label"] == "Su"
            styled(c, bold=True, size=6, color="FFFFFF" if is_sun else "1F4E79",
                   fill="5B9BD5" if is_sun else "DEEBF7")
            c.value = f"{session['label']}\n{session['date'].day}"
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col + offset)].width = 3.1
        col += span

    ws.row_dimensions[head_row].height = 14
    ws.row_dimensions[sub_row].height = 20

    # ---------- task rows grouped by phase ----------
    row = sub_row + 1
    for key, (phase_name, color) in PHASES.items():
        phase_tasks = [t for t in TASKS if t[1] == key]
        if not phase_tasks:
            continue

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        band = ws.cell(row=row, column=1)
        styled(band, bold=True, size=8, color="FFFFFF", fill=color, align="left")
        band.value = f"  {key}  {phase_name}"
        ws.row_dimensions[row].height = 13
        row += 1

        for tid, _, name, t_start, t_end in phase_tasks:
            hours = sum(s["hours"] for s in sessions if t_start <= s["date"] <= t_end)

            styled(ws.cell(row=row, column=1), size=8).value = tid
            styled(ws.cell(row=row, column=2), size=8, align="left").value = name
            styled(ws.cell(row=row, column=3), size=8).value = key
            for offset, value in ((4, t_start), (5, t_end)):
                c = styled(ws.cell(row=row, column=offset), size=8)
                c.value = value
                c.number_format = "mmm dd"
            styled(ws.cell(row=row, column=6), size=8).value = hours

            for idx, session in enumerate(sessions):
                c = styled(ws.cell(row=row, column=first_session_col + idx), size=7)
                if t_start <= session["date"] <= t_end:
                    c.fill = PatternFill("solid", fgColor=color)

            ws.row_dimensions[row].height = 12.5
            row += 1

    # ---------- bottom blocks: legend (left) + project info (right) ----------
    row += 2
    block_top = row

    # legend, columns A-F
    lrow = block_top
    title = styled(ws.cell(row=lrow, column=1), bold=True, size=9, align="left", border=False)
    title.value = "LEGEND (SDLC phases)"
    ws.merge_cells(start_row=lrow, start_column=1, end_row=lrow, end_column=6)
    ws.row_dimensions[lrow].height = 15
    lrow += 1
    for key, (phase_name, color) in PHASES.items():
        chip = styled(ws.cell(row=lrow, column=1), size=8, fill=color,
                      color="FFFFFF", bold=True, align="center")
        chip.value = key
        ws.merge_cells(start_row=lrow, start_column=2, end_row=lrow, end_column=6)
        name = styled(ws.cell(row=lrow, column=2), size=8, align="left", border=False)
        name.value = phase_name
        ws.row_dimensions[lrow].height = 13
        lrow += 1

    # project information, to the right of the legend.
    # Bold only the labels; leave the values in normal weight.
    info_c0, info_c1 = 10, 28

    def bold_label(label, value):
        return CellRichText(
            TextBlock(InlineFont(rFont=FONT, sz=8, b=True), label),
            TextBlock(InlineFont(rFont=FONT, sz=8, b=False), value),
        )

    # (plain_text, bold)  OR  ("split", label, value)
    info_rows = [
        ("Group Members:", True),
        ("Aballa, Gerjhon", False),
        ("Abragan, John Lloyd", False),
        ("Cuenca, John Adam A.", False),
        ("split", "Subject:  ", "CPE 048: Embedded Systems"),
        ("split", "Section:  ", "COC-FA-CPE4-01"),
    ]
    prow = block_top
    for entry in info_rows:
        cell = styled(ws.cell(row=prow, column=info_c0), size=8, align="left", border=False)
        if entry[0] == "split":
            cell.value = bold_label(entry[1], entry[2])
        else:
            text, bold = entry
            cell.font = Font(name=FONT, bold=bold, size=8)
            cell.value = text
        ws.merge_cells(start_row=prow, start_column=info_c0, end_row=prow, end_column=info_c1)
        prow += 1

    row = max(lrow, prow)

    # ---------- print setup: landscape, short bond, one page wide ----------
    ws.freeze_panes = f"{get_column_letter(first_session_col)}{sub_row + 1}"
    ws.print_title_rows = f"{head_row}:{sub_row}"
    ws.print_area = f"A1:{last_letter}{row}"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER      # short bond, 8.5 x 11 in
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.3
    ws.sheet_view.showGridLines = False

    wb.save(OUTFILE)
    print(f"Saved {OUTFILE.name}")
    print(f"  {len(weeks)} weeks, {len(sessions)} sessions, {total_hours:.0f} hours, {len(TASKS)} tasks")
    print(f"  first session {sessions[0]['date']}  last session {sessions[-1]['date']}")


if __name__ == "__main__":
    main()
